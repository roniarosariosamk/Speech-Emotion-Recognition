"""Generate the presentation from the REAL final results (results/final/final_summary.json).

    python summarize_results.py && python build_ppt.py --name "Your Name"
"""
import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

ROOT = Path(__file__).resolve().parent; RES = ROOT / "results" / "final"
ap = argparse.ArgumentParser(); ap.add_argument("--name", default="Ronia")
ap.add_argument("--out", default=str(ROOT.parent / "Speech_Emotion_Recognition_Presentation.pptx")); args = ap.parse_args()
if not (RES / "final_summary.json").is_file():
    raise SystemExit("results/final/final_summary.json not found - run:  python summarize_results.py")
S_ = json.loads((RES / "final_summary.json").read_text()); SYN = False
SEL = S_["selected_final_model"]; T = SEL["test"]; CF = SEL["config"]; SVM = S_["svm_baseline"]; DS = S_["dataset"]; SP = DS["split"]
DONE = [r for r in S_["experiments"] if r["status"] == "completed"]
f4 = lambda x: f"{x:.3f}"
C = lambda h: RGBColor.from_string(h)
INDIGO, CYAN, AMBER, SLATE, MUTED, LIGHT, WHITE, DARK = map(C, ["4F46E5", "06B6D4", "F59E0B", "1E293B", "64748B", "EEF2FF", "FFFFFF", "0F172A"])
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def rect(s, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None):
    r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h)); r.fill.solid(); r.fill.fore_color.rgb = fill
    if line: r.line.color.rgb = line; r.line.width = Pt(1.25)
    else: r.line.fill.background()
    r.shadow.inherit = False; return r


def text(s, x, y, w, h, t, size=16, color=SLATE, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Calibri"):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(t if isinstance(t, list) else [t]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        r = p.add_run(); r.text = line; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = font
    return tb


def bullets(s, x, y, w, items, size=17, gap=10):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.5)); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = Pt(gap)
        r0 = p.add_run(); r0.text = "\u25A0  "; r0.font.size = Pt(size - 6); r0.font.color.rgb = INDIGO
        head, _, rest = it.partition("|")
        if rest:
            r1 = p.add_run(); r1.text = head; r1.font.bold = True; r1.font.size = Pt(size); r1.font.color.rgb = SLATE; r1.font.name = "Calibri"
            r2 = p.add_run(); r2.text = rest; r2.font.size = Pt(size); r2.font.color.rgb = SLATE; r2.font.name = "Calibri"
        else:
            r1 = p.add_run(); r1.text = head; r1.font.size = Pt(size); r1.font.color.rgb = SLATE; r1.font.name = "Calibri"
    return tb


def slide(title, kicker="", notes=""):
    s = prs.slides.add_slide(BLANK); rect(s, 0, 0, 13.333, 0.12, INDIGO)
    if kicker: text(s, 0.7, 0.42, 9, 0.3, kicker.upper(), 12, INDIGO, True)
    text(s, 0.7, 0.72, 11.9, 0.8, title, 32, DARK, True)
    text(s, 0.7, 7.08, 8, 0.25, "Speech Emotion Recognition  |  MFCC + CNN-BiLSTM-Attention", 10, MUTED)
    text(s, 12.0, 7.08, 0.65, 0.25, str(len(prs.slides)), 10, MUTED, align=PP_ALIGN.RIGHT)
    if False:
        b = rect(s, 7.2, 0.38, 5.45, 0.34, C("FEF3C7"), MSO_SHAPE.ROUNDED_RECTANGLE, AMBER)
        text(s, 7.2, 0.38, 5.45, 0.34, "SYNTHETIC SMOKE-TEST DATA - NOT RAVDESS RESULTS", 11, C("92400E"), True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    if notes: s.notes_slide.notes_text_frame.text = notes
    return s


def card(s, x, y, w, h, head, body, accent=INDIGO, bsize=14):
    rect(s, x, y, w, h, LIGHT, MSO_SHAPE.ROUNDED_RECTANGLE); rect(s, x, y + 0.18, 0.08, h - 0.36, accent)
    text(s, x + 0.3, y + 0.2, w - 0.5, 0.4, head, 17, DARK, True); text(s, x + 0.3, y + 0.68, w - 0.5, h - 0.8, body, bsize, SLATE)


def kpi(s, x, y, w, value, label, fill=INDIGO):
    rect(s, x, y, w, 1.55, fill, MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x, y + 0.2, w, 0.8, value, 40, WHITE, True, PP_ALIGN.CENTER); text(s, x, y + 1.0, w, 0.4, label, 13, WHITE, False, PP_ALIGN.CENTER)


def pic(s, name, x, y, w=None, h=None):
    return s.shapes.add_picture(str(RES / name), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


pc = T["per_class"]; tc = SEL["top_confusions"]
low = sorted(pc.items(), key=lambda kv: kv[1]["recall"])[:3]; high = max(pc.items(), key=lambda kv: kv[1]["f1"])
sel_idx = [r["experiment"] for r in DONE].index(SEL["experiment"])
E4 = any(r["experiment"] == "experiment_04" for r in DONE)

# 1 -- title
s = prs.slides.add_slide(BLANK); rect(s, 0, 0, 13.333, 7.5, INDIGO); rect(s, 0, 5.05, 13.333, 0.08, AMBER)
text(s, 0.9, 0.9, 11, 0.4, "INDIVIDUAL MAJOR PROJECT  |  MACHINE LEARNING INTERNSHIP", 14, C("C7D2FE"), True)
text(s, 0.9, 1.5, 11.6, 2.0, ["Emotion Recognition in Speech", "using MFCC Features and Deep Learning"], 42, WHITE, True)
text(s, 0.9, 3.55, 11.5, 0.9, f"CNN-BiLSTM-Attention model  |  RAVDESS  |  evaluated with Macro F1 on unseen actors", 20, C("E0E7FF"))
for i, (v, l) in enumerate([(f4(T["macro_f1"]), "test Macro F1"), (f"{T['accuracy'] * 100:.1f}%", "test accuracy"), (f"{SEL['parameters']:,}", "parameters")]):
    rect(s, 0.9 + i * 2.75, 5.45, 2.5, 1.2, C("4338CA"), MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 0.9 + i * 2.75, 5.55, 2.5, 0.6, v, 30, WHITE, True, PP_ALIGN.CENTER); text(s, 0.9 + i * 2.75, 6.15, 2.5, 0.35, l, 13, C("C7D2FE"), False, PP_ALIGN.CENTER)
text(s, 9.3, 5.55, 3.4, 0.9, [f"Presented by {args.name}", "Machine Learning Internship"], 16, WHITE, False, PP_ALIGN.RIGHT)
s.notes_slide.notes_text_frame.text = "Introduce the project: recognise 8 emotions from speech using MFCC features and a CNN-BiLSTM-Attention network. Headline numbers are on unseen test actors."

# 2 -- problem statement
s = slide("Problem statement", "Why this project", "Emotion is carried by how we speak (pitch, energy, timbre, rhythm), not only by words. The task: classify a speech clip into one of 8 emotions.")
rect(s, 0.7, 1.75, 12, 1.25, LIGHT, MSO_SHAPE.ROUNDED_RECTANGLE); rect(s, 0.7, 1.9, 0.09, 0.95, INDIGO)
text(s, 1.05, 1.9, 11.4, 1.0, "Develop a machine learning system for emotion recognition in speech using MFCC features and a deep learning model, including data collection, preprocessing, model implementation and evaluation using F1-score.", 19, DARK, False, anchor=MSO_ANCHOR.MIDDLE)
card(s, 0.7, 3.3, 3.9, 3.4, "Why it matters", ["Call-centre analytics", "Tutoring and wellbeing tools", "Human-robot interaction", "Accessibility support"], INDIGO, 20)
card(s, 4.75, 3.3, 3.9, 3.4, "Why it is hard", ["Emotion is subjective", "Every speaker sounds different", "Public datasets are small", "Models can memorise voices"], AMBER, 20)
card(s, 8.8, 3.3, 3.9, 3.4, "Our approach", ["MFCC + delta features", "CNN-BiLSTM-Attention network", "Unseen-actor evaluation", "Macro F1 as the main metric"], CYAN, 20)

# 3 -- objectives
s = slide("Objectives", "What we set out to do", "Each objective maps to a module of the code base and is covered by automated tests.")
objs = [("1  Data collection", "RAVDESS speech corpus, labels parsed from filenames"), ("2  Preprocessing", "Resample, trim, normalise, fixed 3 s window, augment (train only)"),
        ("3  MFCC features", "40 MFCC + delta + delta-delta = 120 features x 130 frames"), ("4  Deep model", "CNN -> BiLSTM -> attention -> softmax over 8 emotions"),
        ("5  Baseline", "RBF-SVM on MFCC statistics for fair comparison"), ("6  F1 evaluation", "Macro / weighted / per-class F1 on held-out actors")]
for i, (h, b) in enumerate(objs):
    card(s, 0.7 + (i % 3) * 4.1, 1.8 + (i // 3) * 2.6, 3.9, 2.35, h, b, [INDIGO, CYAN, AMBER][i % 3], 18)

# 4 -- dataset
s = slide("Dataset: RAVDESS speech", "Data collection", "1,440 clips, 24 actors (12F/12M), 2 sentences, 8 emotions. Neutral has a single intensity, so 96 clips; others have 192. Split is by actor to avoid speaker leakage.")
for i, (v, l) in enumerate([(f"{DS['clips']:,}", "clips"), (str(DS["actors"]), "actors"), ("8", "emotions")]): kpi(s, 0.7 + i * 2.3, 1.75, 2.1, v, l, [INDIGO, CYAN, AMBER][i])
text(s, 0.7, 3.55, 7, 0.4, "Actor-independent split", 18, DARK, True)
tot = 12.0 * 0.0 + 6.9; parts = [("Train", SP["train_clips"], SP["train_actors"], INDIGO), ("Val", SP["val_clips"], SP["val_actors"], AMBER), ("Test", SP["test_clips"], SP["test_actors"], CYAN)]
x = 0.7
for n, c, a, col in parts:
    w = tot * c / DS["clips"]; rect(s, x, 4.05, w - 0.04, 0.85, col, MSO_SHAPE.RECTANGLE)
    text(s, x, 4.05, w - 0.04, 0.85, f"{n}: {c}", 14, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(s, x, 4.95, w - 0.04, 0.4, f"actors {a[0]}-{a[-1]}", 12, MUTED, False, PP_ALIGN.CENTER); x += w
bullets(s, 0.7, 5.5, 6.9, ["No actor appears in two partitions (asserted in code)", f"Training grows to {SP['train_after_aug']:,} samples with augmentation; val/test are never augmented"], 14, 6)
pic(s, "class_distribution.png", 7.9, 1.8, w=4.9)
text(s, 7.9, 4.2, 4.9, 1.0, "Labels: neutral, calm, happy, sad, angry, fearful, disgust, surprised (fixed order used everywhere).", 13, MUTED)

# 5 -- methodology
s = slide("Proposed methodology", "Pipeline", "Raw audio -> preprocessing -> MFCC -> normalisation -> CNN-BiLSTM-Attention -> emotion. Each stage is a tested Python module.")
pic(s, "diagram_pipeline.png", 0.7, 1.7, w=11.9)
card(s, 0.7, 4.3, 3.9, 2.5, "No data leakage", "Normalisation statistics come from training data only; augmentation only touches training clips.", INDIGO, 16)
card(s, 4.75, 4.3, 3.9, 2.5, "Unseen-actor testing", "Test actors 21-24 never appear in training or tuning, so scores reflect new voices.", CYAN, 16)
card(s, 8.8, 4.3, 3.9, 2.5, "Validation-based selection", "Checkpoint, early stopping and final-model choice use validation Macro F1, not the test set.", AMBER, 16)

# 6 -- preprocessing + MFCC
s = slide("Preprocessing and MFCC extraction", "Features", "MFCC describe the voice timbre per frame; delta features add how it changes; delta-delta adds how that change accelerates.")
bullets(s, 0.7, 1.75, 6.3, ["Resample|  to 22,050 Hz mono", "Trim silence|  (30 dB) and peak-normalise", "Fixed window|  of 3 s (crop / zero-pad)", "MFCC|  40 coefficients, FFT 2048, hop 512, 128 mel bands",
                            "Delta + delta-delta|  adds speed and acceleration of each coefficient", "Tensor|  130 frames x 120 features per clip", "Standardise|  with training mean / std saved to norm.npz"], 16, 9)
pic(s, "mfcc_example.png", 7.2, 1.8, w=5.6)
card(s, 7.2, 4.65, 5.6, 2.1, "Augmentation (train only)", "2 extra copies per clip: noise, pitch +/-2 semitones, time stretch 0.9-1.1x or time shift.", AMBER, 16)

# 7 -- architecture
s = slide("Model architecture", "CNN + BiLSTM + attention", f"Compact network with {SEL['parameters']:,} parameters. CNN finds local spectral patterns, BiLSTM models context both ways, attention weights the most informative frames.")
pic(s, "diagram_architecture.png", 0.7, 1.65, w=11.9)
card(s, 0.7, 4.35, 3.9, 2.45, "CNN", f"2 Conv1D blocks ({CF['conv1']['filters']} and {CF['conv2']['filters']} filters) with batch norm, ReLU, pooling and dropout.", CYAN, 16)
card(s, 4.75, 4.35, 3.9, 2.45, "BiLSTM", f"{CF['lstm']['units']} units per direction read the sequence forwards and backwards.", INDIGO, 16)
card(s, 8.8, 4.35, 3.9, 2.45, "Attention + softmax", f"Learned frame weights give one utterance vector; Dense({CF['dense']['units']}) and softmax output 8 probabilities.", AMBER, 16)

# 8 -- training and evaluation
s = slide("Training and evaluation", "Setup", "Class weights compensate for fewer neutral clips. The best checkpoint is the epoch with the highest validation Macro F1.")
rows = [("Optimiser", "Adam, lr 1e-3, halved on plateau"), ("Loss", "Cross-entropy + class weights"), ("Batch / epochs", "32 / up to 60"), ("Early stopping", "12 epochs on val Macro F1"),
        ("Best epoch", f"{SEL['best_epoch']} of {SEL['epochs_run']} run"), ("Seed / time", f"{DS['seed']} / {SEL['training_seconds']:.0f} s CPU"), ("Metrics", "Macro F1 (main), weighted F1, per-class, confusion matrix")]
for i, (k, v) in enumerate(rows):
    y = 1.8 + i * 0.68; rect(s, 0.7, y, 5.9, 0.6, LIGHT if i % 2 == 0 else WHITE)
    text(s, 0.85, y, 2.1, 0.6, k, 13, INDIGO, True, anchor=MSO_ANCHOR.MIDDLE); text(s, 2.9, y, 3.7, 0.6, v, 13, SLATE, False, anchor=MSO_ANCHOR.MIDDLE)
pic(s, "training_curves.png", 6.9, 1.8, w=5.9)
text(s, 6.9, 4.35, 5.9, 1.2, "Validation loss stops improving after about epoch 12; the checkpoint from the best validation Macro F1 epoch is kept, which limits over-fitting.", 14, SLATE)

# 9 -- results
s = slide("Results and baseline comparison", "Results", "Test set = 4 unseen actors (240 clips). Validation Macro F1 was used to choose the final model. " + ("Experiment 4 (stronger dropout) was trained once on the real data and is included in the comparison." if E4 else "Experiment 4 (stronger dropout) is implemented but was not run on real data, so no result is claimed for it."))
cd = CategoryChartData(); cd.categories = [r["label"].replace("Experiment ", "Exp ") for r in DONE] + ["SVM"]
cd.add_series("Validation Macro F1", [round(r["val_macro_f1"], 4) for r in DONE] + [None]); cd.add_series("Test Macro F1", [round(r["test_macro_f1"], 4) for r in DONE] + [round(SVM["test_macro_f1"], 4)])
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.7), Inches(1.75), Inches(7.4), Inches(4.9), cd).chart
ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(12)
ch.value_axis.maximum_scale = 0.7; ch.value_axis.minimum_scale = 0; ch.value_axis.tick_labels.font.size = Pt(11); ch.category_axis.tick_labels.font.size = Pt(12); ch.value_axis.has_major_gridlines = False
pl = ch.plots[0]; pl.gap_width = 60; pl.has_data_labels = True; pl.data_labels.number_format = "0.000"; pl.data_labels.number_format_is_linked = False; pl.data_labels.font.size = Pt(11); pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
pl.series[0].format.fill.solid(); pl.series[0].format.fill.fore_color.rgb = AMBER; pl.series[1].format.fill.solid(); pl.series[1].format.fill.fore_color.rgb = INDIGO
kpi(s, 8.4, 1.75, 2.1, f4(T["macro_f1"]), f"{SEL['label']} Macro F1"); kpi(s, 10.6, 1.75, 2.1, f"{T['accuracy'] * 100:.1f}%", "test accuracy", CYAN)
kpi(s, 8.4, 3.45, 2.1, f4(SVM["test_macro_f1"]), "SVM baseline Macro F1", C("64748B")); kpi(s, 10.6, 3.45, 2.1, f"{SEL['gain_over_svm_macro_f1']:+.3f}", "Macro F1 vs SVM", AMBER)
text(s, 8.4, 5.2, 4.3, 1.5, [f"Weighted F1 {f4(T['weighted_f1'])}  |  macro precision {f4(T['macro_precision'])}  |  macro recall {f4(T['macro_recall'])}", "Small test set: differences of a few points between experiments are not conclusive."], 12, MUTED)

# 10 -- confusion matrix and error analysis
s = slide("Confusion matrix and error analysis", "Results", "Observed errors (counts from the test confusion matrix). Explanations are hypotheses, not proven.")
pic(s, "final_confusion_matrix.png", 0.6, 1.65, h=5.3)
text(s, 6.9, 1.7, 5.8, 0.4, "What the model gets right and wrong", 18, DARK, True)
bullets(s, 6.9, 2.2, 5.9, [f"Best:|  {high[0]} (F1 {f4(high[1]['f1'])}), angry recall {f4(pc['angry']['recall'])}", "Lowest recall:|  " + ", ".join(f"{k} {f4(v['recall'])}" for k, v in low),
                           f"Neutral over-predicted:|  precision only {f4(pc['neutral']['precision'])}"] + [f"{c['true']} -> {c['predicted']}|  {c['count']} clips ({c['share_of_true_class']:.0%})" for c in tc[:4]], 14, 7)
text(s, 6.9, 6.0, 5.8, 0.9, "Possible reasons (untested): similar low-energy delivery in sad / calm / neutral, shared high arousal between fearful / angry / happy, few training actors.", 12, MUTED)

# 11 -- limitations and future work
s = slide("Limitations and future work", "Honest assessment", "State clearly: acted speech, 4 test actors, one split. The model must not be presented as understanding human emotion in general.")
card(s, 0.7, 1.75, 6.0, 5.0, "Limitations", ["Small, acted dataset (24 actors, 2 sentences)", "Only 4 test actors: noisy estimate", "Neutral class has half the clips", "Experiments 1-3 compared with test scores visible: mildly optimistic", "No noise / cross-corpus / fairness tests", ("Experiment 4 is a single run (one seed)" if E4 else "Experiment 4 not executed on real data")], AMBER, 19)
card(s, 6.9, 1.75, 5.8, 5.0, "Future work", [("Repeat Exp 3 vs 4 with several seeds" if E4 else "Run Experiment 4 and compare fairly"), "Actor-grouped cross-validation", "SpecAugment, mixup, more corpora (TESS, CREMA-D)", "Prosodic features / pre-trained speech encoders", "Real-time demo with privacy safeguards"], CYAN, 19)

# 12 -- conclusion
s = slide("Conclusion", "Summary", "Complete pipeline: data -> MFCC -> CNN-BiLSTM-Attention -> F1 evaluation, with tests, inference script and documentation.")
bullets(s, 0.7, 1.8, 7.4, [f"Delivered|  every requirement: data, preprocessing, MFCC, deep model, F1 evaluation", f"Selected model|  {SEL['label']} ({SEL['parameters']:,} parameters), chosen by validation Macro F1",
                           f"Test result|  Macro F1 {f4(T['macro_f1'])}, accuracy {T['accuracy'] * 100:.1f}% on unseen actors", f"Baseline|  beats the tuned SVM ({f4(SVM['test_macro_f1'])}) by {SEL['gain_over_svm_macro_f1']:+.3f} Macro F1",
                           "Honest view|  uneven across emotions; sad and fearful remain hard", "Reproducible|  seeds, tests, one-command evaluation and prediction"], 17, 12)
rect(s, 8.6, 1.8, 4.1, 4.6, INDIGO, MSO_SHAPE.ROUNDED_RECTANGLE); text(s, 8.6, 2.5, 4.1, 0.9, "Thank you", 36, WHITE, True, PP_ALIGN.CENTER); text(s, 8.6, 3.6, 4.1, 0.6, "Questions?", 22, C("C7D2FE"), False, PP_ALIGN.CENTER)
text(s, 8.8, 4.7, 3.7, 1.2, ["python predict.py clip.wav", "python summarize_results.py"], 13, C("E0E7FF"), False, PP_ALIGN.CENTER, font="Consolas")

prs.save(args.out); print("wrote", args.out, "-", len(prs.slides), "slides")
